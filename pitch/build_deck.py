"""Build the Sehat pitch deck (pitch/Sehat-GRAVITON-2026.pptx) from real screenshots in pitch/shots.
Simple, patient-first: one idea per slide, big type, a patient story (Rahul and his mother Sunita) running through it."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
SHOT = lambda n: os.path.join(HERE, 'shots', n + '.png')
BLUE, LIGHT, INK, MUTED, WHITE, TINT, GREEN, LINE = (RGBColor(0x07, 0x74, 0xC4), RGBColor(0x8E, 0xBE, 0xF9), RGBColor(0x0F, 0x17, 0x2A), RGBColor(0x5B, 0x6B, 0x80),
                                                   RGBColor(255, 255, 255), RGBColor(0xEA, 0xF3, 0xFE), RGBColor(0x05, 0x96, 0x69), RGBColor(0xDC, 0xE7, 0xF5))
DISPLAY, BODY = 'Manrope', 'Plus Jakarta Sans'
LINK = 'https://sehat-ashen.vercel.app'
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
N = [0]


def box(s, x, y, w, h, fill=TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); b.fill.solid(); b.fill.fore_color.rgb = fill
    if line: b.line.color.rgb = line; b.line.width = Pt(1.25)
    else: b.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0] = 0.08
    b.shadow.inherit = False
    return b


def text(s, x, y, w, h, lines, size=20, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT, gap=8, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(gap)
        parts = line if isinstance(line, tuple) else ((line, bold, color),)
        for t, b_, c_ in parts:
            r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = b_; r.font.color.rgb = c_; r.font.name = font
    return tb


def slide(title, kicker=None):
    N[0] += 1; s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, 13.333, 0.1, BLUE, MSO_SHAPE.RECTANGLE)
    if kicker: text(s, 0.7, 0.42, 9, 0.4, kicker.upper(), 13, True, BLUE)
    text(s, 0.7, 0.75, 12, 0.9, title, 36, True, INK, DISPLAY)
    text(s, 0.7, 6.95, 9, 0.35, 'Sehat  ·  Team ANTIJAREX  ·  GRAVITON 2026', 11, False, MUTED)
    text(s, 12.0, 6.95, 0.7, 0.35, str(N[0]), 11, True, MUTED, align=PP_ALIGN.RIGHT)
    return s


def shot(s, name, x, y, w):
    if not os.path.exists(SHOT(name)): return
    from PIL import Image
    iw, ih = Image.open(SHOT(name)).size; h = w * ih / iw
    box(s, x - 0.08, y - 0.08, w + 0.16, h + 0.16, WHITE, line=LINE)
    s.shapes.add_picture(SHOT(name), Inches(x), Inches(y), width=Inches(w))


def stat(s, x, y, w, big, small, src=None, fill=TINT):
    box(s, x, y, w, 2.5, fill)
    text(s, x + 0.3, y + 0.25, w - 0.6, 0.9, big, 40, True, BLUE, DISPLAY)
    text(s, x + 0.3, y + 1.15, w - 0.6, 1.0, small, 16, False, INK)
    if src: text(s, x + 0.3, y + 2.05, w - 0.6, 0.4, src, 10, False, MUTED)


def points(s, x, y, w, items, size=20, icon_color=BLUE):
    for i, (ic, head, sub) in enumerate(items):
        yy = y + i * 1.05
        c = box(s, x, yy, 0.62, 0.62, TINT, MSO_SHAPE.OVAL); c.text_frame.text = ic; c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        c.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        text(s, x + 0.85, yy - 0.05, w - 0.85, 0.45, head, size, True, INK)
        if sub: text(s, x + 0.85, yy + 0.38, w - 0.85, 0.5, sub, 14, False, MUTED)


# 1 · title
N[0] += 1; s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, TINT, MSO_SHAPE.RECTANGLE)
m = box(s, 0.7, 1.35, 0.9, 0.9, BLUE); m.text_frame.text = '+'; m.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER; r = m.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(44); r.font.bold = True; r.font.color.rgb = WHITE
text(s, 1.8, 1.3, 5, 1, 'Sehat', 54, True, INK, DISPLAY)
text(s, 0.7, 2.55, 6.2, 1.6, 'Your health, connected:\nfrom doctor to doorstep.', 34, True, INK, DISPLAY, gap=2)
text(s, 0.7, 4.25, 5.9, 1.2, 'One app for patients and families, live-connected to hospitals, doctors, pharmacies and suppliers. In your language.', 18, False, MUTED)
text(s, 0.7, 5.55, 6, 0.4, 'GRAVITON 2026  ·  Team ANTIJAREX  ·  CGC University', 14, True, BLUE)
btn = box(s, 0.7, 6.15, 4.6, 0.7, BLUE); btn.adjustments[0] = 0.5
tf = btn.text_frame; tf.text = '🔗  Try it live: sehat-ashen.vercel.app'; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.runs[0]; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = BODY; btn.click_action.hyperlink.address = LINK  # whole button is the link, so the text keeps its white style
shot(s, 'home', 7.0, 1.4, 5.7)

# 2 · problem, told for the patient
s = slide('Getting care in India is a maze for patients', 'The problem')
stat(s, 0.7, 1.9, 3.85, '1 : 1,194', 'One allopathic doctor for every ~1,200 people. Long queues, rushed visits.', 'Allopathic-only estimate (govt 1:834 includes AYUSH)')
stat(s, 4.75, 1.9, 3.85, '2 in 3', 'Chronic patients in one study were not taking medicines as prescribed.', 'Chennai urban health centre study (WHO IMSEAR)')
stat(s, 8.8, 1.9, 3.85, '₹1.61 L cr', 'Spent by Indians on medicines in a year, mostly from their own pocket.', 'National Health Accounts 2022-23')
text(s, 0.7, 4.75, 12, 1.6, ['Patients run between clinic, pharmacy and lab, lose prescriptions, cannot read their reports,',
                            'and elders forget doses. Nothing is connected.'], 20, False, INK, gap=2)

# 3 · patient story
s = slide("Meet Rahul and his mother Sunita", 'A day without vs with Sehat')
box(s, 0.7, 1.85, 5.85, 4.75, RGBColor(0xFE, 0xF2, 0xF2))
text(s, 1.0, 2.05, 5.3, 0.5, 'Without Sehat', 22, True, RGBColor(0xB9, 0x1C, 0x1C))
text(s, 1.0, 2.7, 5.3, 3.8, ['✗  Waits hours to see a doctor', '✗  Carries a paper prescription to 3 chemists', '✗  Cannot understand his lab report',
                            '✗  Mother forgets her BP tablet', '✗  No one knows if she needs help'], 18, gap=10)
box(s, 6.8, 1.85, 5.85, 4.75, RGBColor(0xEC, 0xFD, 0xF5))
text(s, 7.1, 2.05, 5.3, 0.5, 'With Sehat', 22, True, GREEN)
text(s, 7.1, 2.7, 5.3, 3.8, ['✓  Books the right doctor in a minute', '✓  Prescription goes straight to his pharmacy', '✓  Report explained in Hindi or Punjabi',
                            '✓  Mother gets spoken reminders', '✓  One tap SOS alerts family and hospital'], 18, gap=10)

# 4 · what it is
s = slide('One app for patients, connected to everyone', 'The solution')
points(s, 0.7, 1.95, 6.2, [('🙂', 'Patient & family', 'Book, track medicines, understand reports'), ('🏥', 'Hospital', 'Sees every booking live, beds, doctors'),
                         ('🩺', 'Doctor', 'Queue by urgency, prescription in one tap'), ('💊', 'Pharmacy', 'Gets prescriptions live, sends the bill'),
                         ('🚚', 'Supplier', 'Restocks pharmacies before they run out')], 19)
shot(s, 'login', 7.2, 1.95, 5.5)

# 5 · the journey
s = slide('From symptom to doorstep, live', "Rahul's journey")
steps = [('1', 'Book', 'Types "fever": right doctor and earliest slot'), ('2', 'Visit', 'Hospital checks him in, doctor signs the prescription'),
         ('3', 'Approve', 'His pharmacy sends the bill, Rahul approves'), ('4', 'Delivered', 'Rider brings it home, tracked live')]
for i, (n, h, sub) in enumerate(steps):
    x = 0.7 + i * 3.05; box(s, x, 1.85, 2.85, 1.75)
    c = box(s, x + 0.25, 2.05, 0.55, 0.55, BLUE, MSO_SHAPE.OVAL); c.text_frame.text = n; p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.runs[0].font.size = Pt(18); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = WHITE
    text(s, x + 0.95, 2.05, 1.8, 0.5, h, 22, True, INK, DISPLAY); text(s, x + 0.25, 2.75, 2.45, 0.85, sub, 13, False, MUTED)
shot(s, 'patient', 1.45, 3.85, 4.6); shot(s, 'pharmacy', 7.25, 3.85, 4.6)

# 6 · AI in your language
s = slide('Your report, explained in your language', 'AI that helps patients')
points(s, 0.7, 1.95, 5.4, [('📄', 'Photo or PDF of any report', 'Lab reports, prescriptions, discharge summaries'), ('🗣️', '10 Indian languages', 'Hindi, Punjabi, Tamil, Bengali, Telugu and more'),
                         ('🎙️', 'Dr. Saathi voice assistant', 'Ask by speaking, answers in the same language'), ('🛡️', 'Safe by design', 'Explains, never diagnoses; our code flags values')], 19)
shot(s, 'records', 6.4, 1.95, 6.3)

# 7 · senior mode
s = slide('Senior Mode: built for Dadi and Nani', 'For elders')
points(s, 0.7, 1.95, 5.4, [('🔘', 'Three big buttons, nothing else', 'Medicines now · talk to Dr. Saathi · SOS'), ('🔊', 'Spoken reminders', 'In her language, with a lock against double doses'),
                         ('🔔', 'Family is alerted', 'If a dose is 45 minutes late, her son is told'), ('🆘', 'One-tap SOS', 'Alerts the hospital with her history, calls 112')], 19)
shot(s, 'senior', 6.4, 1.95, 6.3)

# 8 · journal + privacy
s = slide('Private by default. Shared only with consent.', 'Trust')
points(s, 0.7, 1.95, 5.4, [('📔', 'Daily journal', 'Mood, symptoms, sleep in a minute; a pre-visit summary'), ('🔒', 'Patient owns it', 'Pharmacy and others never see the journal'),
                         ('✅', 'Doctor sees it only when shared', 'Stop sharing any time'), ('👁️', 'Each role sees only its work', 'Hospital: its bookings · pharmacy: its orders')], 19)
shot(s, 'journal', 6.4, 1.95, 6.3)

# 9 · pricing, kept simple
s = slide('Free for patients. Simple monthly plans for providers.', 'Pricing')
box(s, 0.7, 1.9, 4.6, 4.6, BLUE)
text(s, 1.0, 2.2, 4.0, 0.5, 'PATIENTS & FAMILIES', 14, True, WHITE)
text(s, 1.0, 2.7, 4.0, 1.2, '₹0', 72, True, WHITE, DISPLAY)
text(s, 1.0, 4.0, 4.0, 2.3, ['Free forever', 'Booking, tracking, reports explained, journal, Senior Mode'], 17, False, WHITE, gap=6)
rows = [('🩺  Doctor', '₹999 / month'), ('🧑‍⚕️  Nurse', '₹499 / month'), ('💊  Pharmacy', '₹1,499 / month'), ('🚚  Supplier', '₹2,499 / month'), ('🏥  Hospital', '₹4,999 / month')]
for i, (k, v) in enumerate(rows):
    y = 1.9 + i * 0.84; box(s, 5.6, y, 7.05, 0.72, TINT if i % 2 == 0 else WHITE, line=None if i % 2 == 0 else LINE)
    text(s, 5.9, y + 0.13, 3.6, 0.5, k, 18, True, INK); text(s, 9.3, y + 0.13, 3.2, 0.5, v, 18, True, BLUE, align=PP_ALIGN.RIGHT)
text(s, 5.6, 6.15, 7.0, 0.5, 'Providers pay because Sehat brings them patients and saves staff time. Proposed launch pricing.', 12, False, MUTED)

# 10 · proof
s = slide('Working today, not a mock-up', 'Proof')
stat(s, 0.7, 1.9, 2.9, '6 roles', 'Live across separate devices, synced in about a second')
stat(s, 3.8, 1.9, 2.9, '140/140', 'Fields read correctly on our printed report benchmark')
stat(s, 6.9, 1.9, 2.9, '10', 'Indian languages for reports and voice; 8 in Senior Mode')
stat(s, 10.0, 1.9, 2.65, '0', 'Server libraries to install: pure Python, deployed on Vercel')
text(s, 0.7, 4.8, 12, 1.4, ['Try it now: sehat-ashen.vercel.app  →  Live demo  →  ▶ Auto-play', 'Benchmark is on clean printed samples; real handwritten prescriptions go through a review screen.'], 16, False, MUTED, gap=4)

# 11 · why we win
s = slide('Why Sehat wins', 'Our edge')
points(s, 0.7, 1.95, 12, [('🔗', 'Connected, not another single app', 'Patient, hospital, doctor, pharmacy and supplier on one live network'),
                        ('🗣️', 'Indian languages, end to end', 'Reports, chat, voice and Senior Mode'), ('🔒', 'Consent built in', 'Patients decide who sees what'),
                        ('🧓', 'Designed for elders too', 'The Dadi / Nani test: big buttons, voice, SOS'), ('🧱', 'Honest foundation', 'Sehat Copilot was our earlier patient app; at GRAVITON we built the network around it')], 19)

out = os.path.join(HERE, 'Sehat-GRAVITON-2026.pptx'); prs.save(out); print(out)
